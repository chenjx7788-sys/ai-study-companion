# -*- coding: utf-8 -*-
"""P0 验证：去噪管道（4 规则）+ 结构化切分 + DOCX page_no + PDF section_path

运行：backend 目录下  python test_p0_rag.py
全部断言通过则输出 ALL PASS。
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from app.services import cleaner, vector, parser   # noqa: E402


def sep(title):
    print(f"\n===== {title} =====")


# ---------- 1. 去噪：页眉页脚 + 水印 + 页码 + 乱码 + 重复 ----------
sep("1. 去噪管道")
pages = []
for i in range(1, 7):   # 6 页 PDF 模拟块
    pages.append({
        "content": (
            "XX公司内部培训资料\n"                     # 页眉（每页重复）
            f"第 {i} 章 需求管理的核心概念与实践方法。\n"
            "本页正文内容讨论用户价值与交易模型的关系。\n"
            "内部资料 注意保密\n"                        # 页中水印（<30字 无句末标点 每页重复）
            "仅供内部使用\x00\uf020乱码行测试\ufffd\ufffd\n"   # 乱码行（含 PUA/控制符/替换符）
            "机密文件 请勿外传\n"                      # 尾行水印（命中页脚规则）
            f"- {i} -"                                # 页码
        ),
        "page_no": i,
        "section_path": "",
    })
# 重复块：两块内容仅空白不同
pages.append({"content": "第 1 章 需求管理的核心概念与实践方法。\n本页正文内容讨论用户价值与交易模型的关系。",
              "page_no": 7, "section_path": ""})
pages.append({"content": "第 1 章 需求管理的核心概念与实践方法。 本页正文内容讨论用户价值与交易模型的关系。",
              "page_no": 8, "section_path": ""})

cleaned, report = cleaner.clean_chunks(pages, "pdf")
print("报告:", report)
assert report["header_footer_lines"] >= 12, "页眉+尾行水印应被剔除"
assert report["watermark_lines"] == 6, "6 个页中水印行应被剔除"
assert report["garbled_lines"] == 6, "6 个乱码行应被剔除"
assert report["page_no_lines"] == 6, "6 个页码行应被剔除"
assert report["dup_blocks"] == 2, "页7/8 归一化后与页1重复，应去重 2 个"
assert len(cleaned) == 6, f"应剩 6 块，实际 {len(cleaned)}"
body = cleaned[0]["content"]
assert "XX公司内部培训资料" not in body and "机密文件" not in body and "内部资料" not in body
assert "\ufffd" not in body and "- 1 -" not in body
assert "需求管理" in body, "正文必须保留"
assert "用户价值与交易模型的关系。" in body, "带句号的重复正文句不得被水印规则误伤"
print("正文保留 ✓ 页眉/水印/乱码/页码/重复块全部剔除 ✓")

# 少于 4 页不做跨页统计（防误伤）
few = [{"content": "标题\n正文内容只有一页\n标题", "page_no": 1, "section_path": ""}]
c2, r2 = cleaner.clean_chunks(few, "pdf")
assert c2[0]["content"] == few[0]["content"] and r2["watermark_lines"] == 0
print("少页文档不误伤 ✓")

# 非分页格式（md/docx/media）不做页眉判定但仍去乱码+去重
md_chunks = [{"content": "标题\n正文一", "page_no": 1, "section_path": ""},
             {"content": "标题\n正文一", "page_no": 2, "section_path": ""}]
c3, r3 = cleaner.clean_chunks(md_chunks, "md")
assert len(c3) == 1 and r3["dup_blocks"] == 1 and r3["watermark_lines"] == 0
print("md 只去重、不做页眉判定 ✓")


# ---------- 2. 结构化切分 ----------
sep("2. 结构化切分")

# 2a. 短文本原样返回
assert vector.split_text("短文本", size=700) == ["短文本"]

# 2b. 句子不被拦腰切断：每块结尾应是句边界
sentences = [f"这是第{i}个完整的句子，用来测试切分是否尊重句子的边界。" for i in range(40)]
text = "\n".join(sentences)
blocks = vector.split_text(text, size=200, overlap=30)
assert len(blocks) > 1
for b in blocks:
    assert len(b) <= 200, f"块超长: {len(b)}"
    lines = b.split("\n")
    # 除重叠携带外，每行应是完整句子（以句号结尾）
    for ln in lines:
        assert ln.endswith("。"), f"句子被切断: {ln!r}"
print(f"句边界保持 ✓（{len(sentences)} 句 → {len(blocks)} 块，无断句）")

# 2c. 句级重叠：相邻块存在内容衔接
if len(blocks) >= 2:
    tail_of_prev = blocks[0].split("\n")[-1]
    assert tail_of_prev in blocks[1] or blocks[1].split("\n")[0] in blocks[0], "应有句级重叠"
    print(f"句级重叠 ✓（前块末句 → 后块开头，≤30 字）")

# 2d. 段落优先：空行分段不拆开
paras = "\n\n".join([f"段落{i}。" + "内容。" * 10 for i in range(10)])
blocks2 = vector.split_text(paras, size=120, overlap=0)
for b in blocks2:
    first = b.split("。")[0]
    assert first.startswith("段落"), f"段落被拆: {first!r}"
print("段落边界优先 ✓")

# 2e. 列表不拆：连续列表行整体打包（空行分段内）
list_block = "要点列表：\n" + "\n".join(f"- 第{i}个要点内容" for i in range(8))
blocks3 = vector.split_text(list_block, size=300, overlap=0)
assert len(blocks3) == 1 and blocks3[0].count("- 第") == 8
print("列表整块保留 ✓")

# 2f. 滑窗兜底：单句 2000 字无标点
monster = "字" * 2000
blocks4 = vector.split_text(monster, size=700, overlap=100)
assert len(blocks4) >= 3 and all(len(b) <= 700 for b in blocks4)
print(f"滑窗兜底 ✓（2000 字无标点 → {len(blocks4)} 块）")

# 2g. 默认参数（settings.chunk_size=700）正常工作
big = "\n\n".join(["段落。" * 200 for _ in range(5)])
b5 = vector.split_text(big)
assert all(len(b) <= 700 for b in b5)
print(f"默认 size=700 ✓（{len(b5)} 块）")


# ---------- 3. DOCX page_no + PDF section_path（合成文件） ----------
sep("3. 解析器修复")
import tempfile, os
from docx import Document

tmp = Path(tempfile.mkdtemp())

# 3a. DOCX 无分页符 → 按 ~700 字/页估算
doc = Document()
doc.add_heading("第一章 引言", level=1)
for i in range(30):
    doc.add_paragraph("这是一个测试段落，包含足够多的文字来验证页码估算逻辑是否正常工作。" * 2)
docx_path = tmp / "t1.docx"
doc.save(str(docx_path))
chunks = parser.parse_docx(str(docx_path))
pages_seen = sorted({c["page_no"] for c in chunks})
assert len(pages_seen) > 1, f"page_no 仍恒 1: {pages_seen}"
assert chunks[0]["section_path"] == "第一章 引言"
print(f"DOCX 估算页码 ✓（页码分布 {pages_seen}，section_path 正确）")

# 3b. DOCX 有显式分页符 → 按分页符翻页
doc2 = Document()
doc2.add_paragraph("第一页内容。")
p = doc2.add_paragraph()
run = p.add_run()
from docx.oxml.ns import qn
br = run._r.makeelement(qn("w:br"), {qn("w:type"): "page"})
run._r.append(br)
doc2.add_paragraph("第二页内容。")
docx2_path = tmp / "t2.docx"
doc2.save(str(docx2_path))
chunks2 = parser.parse_docx(str(docx2_path))
pg = {c["content"][:3]: c["page_no"] for c in chunks2}
assert pg.get("第一页") == 1 and pg.get("第二页") == 2, f"分页符翻页失败: {pg}"
print("DOCX 分页符翻页 ✓（第一页→1，第二页→2）")

# 3c. PDF outline → section_path（reportlab 不一定有，用 pypdf 合成）
from pypdf import PdfWriter
writer = PdfWriter()
for _ in range(4):
    writer.add_blank_page(width=612, height=792)
writer.add_outline_item("第一章 概念", 0)
writer.add_outline_item("1.1 定义", 1, parent=writer.outline[-1] if hasattr(writer, "outline") else None)
pdf_path = tmp / "t3.pdf"
with open(pdf_path, "wb") as f:
    writer.write(f)
from pypdf import PdfReader
reader = PdfReader(str(pdf_path))
mpages, mpaths = parser._pdf_section_marks(reader)
assert mpages and "第一章 概念" in mpaths[0], f"outline 解析失败: {list(zip(mpages, mpaths))}"
print(f"PDF 书签映射 ✓（{list(zip(mpages, mpaths))}）")

os.system(f'rd /s /q "{tmp}" 2>nul')


# ---------- 4. 配置化（切分参数 + 清洗开关走设置页存储） ----------
sep("4. 切分与清洗规则配置化")
from app.services import settings_store

orig_path = settings_store.STORE_PATH
tmp_store = Path(tempfile.mkdtemp()) / "llm_settings.json"
settings_store.STORE_PATH = tmp_store   # 隔离：不碰真实配置文件
try:
    # 4a. 默认值来自 config（无配置文件时）
    conf = settings_store.load()
    assert conf["chunk_size"] == 700 and conf["clean_dedup"] is True

    # 4b. 保存自定义切分参数 → split_text 自动读取
    settings_store.save({"chunk_size": 400, "chunk_overlap": 40,
                         "clean_watermark": False, "clean_dedup": False})
    big2 = "\n\n".join(["段落内容。" * 60 for _ in range(4)])   # 每段 300 字
    bs = vector.split_text(big2)
    assert all(len(b) <= 400 for b in bs), f"应按 400 字切分: {[len(b) for b in bs]}"
    assert vector._chunking_params() == (400, 40)
    print("split_text 读取用户配置 ✓（400/40 生效）")

    # 4c. 钳制：非法值回安全范围
    settings_store.save({"chunk_size": 99999, "chunk_overlap": -5})
    assert vector._chunking_params() == (2000, 0)
    settings_store.save({"chunk_size": 400, "chunk_overlap": 40,
                         "clean_watermark": False, "clean_dedup": False})
    print("参数钳制 ✓（99999→2000，-5→0）")

    # 4d. 关闭水印+去重开关 → 对应规则跳过
    pages2 = []
    for i in range(1, 6):
        pages2.append({"content": f"页眉ABC\n正文第{i}段内容。\n机密\n12", "page_no": i, "section_path": ""})
    pages2.append({"content": "页眉ABC\n正文第1段内容。\n机密\n12", "page_no": 6, "section_path": ""})
    c4, r4 = cleaner.clean_chunks(pages2, "pdf")
    assert r4["watermark_lines"] == 0 and r4["dup_blocks"] == 0 and len(c4) == 6
    assert r4["header_footer_lines"] >= 5 and r4["page_no_lines"] == 6   # 其余规则仍生效
    print("开关关闭水印/去重 ✓（页眉页脚与页码仍生效）")

    # 4e. 显式 opts 覆盖（供测试/调用方指定）
    c5, r5 = cleaner.clean_chunks(pages2, "pdf", opts={"header_footer": False, "watermark": False,
                                                       "garbled": False, "dedup": False})
    assert all(v == 0 for v in r5.values()) and len(c5) == 6
    print("显式 opts 全关 ✓（原文原样通过）")
finally:
    settings_store.STORE_PATH = orig_path


sep("ALL PASS")
print("P0 全部验证通过")
